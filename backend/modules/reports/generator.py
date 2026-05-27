"""
Excel Report Generator
Produces a professional, colour-coded .xlsx reconciliation report.
"""

import os
import uuid
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils  import get_column_letter

# ── Colour palette ────────────────────────────────────────────────────
_C = {
    "header_bg"  : "1F4E79",
    "matched"    : "C6EFCE",
    "missing"    : "FFC7CE",
    "discrepancy": "FFEB9C",
    "duplicate"  : "FCE4D6",
    "white"      : "FFFFFF",
}

_HEADER_FONT  = Font(bold=True, color=_C["white"], size=11, name="Calibri")
_HEADER_FILL  = PatternFill("solid", fgColor=_C["header_bg"])
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CELL_ALIGN   = Alignment(vertical="center", wrap_text=False)
_THIN         = Side(style="thin", color="D0D0D0")
_BORDER       = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOLD         = Font(bold=True, name="Calibri", size=10)
_NORMAL       = Font(name="Calibri", size=10)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _write_headers(ws, headers: list, row: int = 1):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border    = _BORDER
    ws.row_dimensions[row].height = 22


def _write_data_row(ws, values: list, row: int, fill: PatternFill = None):
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.font      = _NORMAL
        cell.alignment = _CELL_ALIGN
        cell.border    = _BORDER
        if fill:
            cell.fill = fill


def _autofit(ws):
    for col in ws.columns:
        letter  = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 4, 14), 55)


def _add_sheet(wb, title: str, headers: list, rows: list, row_fn, fill_color: str):
    ws = wb.create_sheet(title=title)
    _write_headers(ws, headers)
    row_fill = _fill(fill_color) if fill_color else None
    for r, item in enumerate(rows, 2):
        _write_data_row(ws, row_fn(item), r, fill=row_fill)
    if not rows:
        ws.cell(row=2, column=1, value="No records found").font = Font(italic=True, color="888888")
    _autofit(ws)
    ws.freeze_panes = "A2"
    return ws


def generate_excel_report(result: dict, output_dir: str) -> str:
    """
    Build a full colour-coded Excel report from reconciliation result.
    Returns the file path of the saved .xlsx file.
    """
    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = "📋 Summary"

    # ── Summary sheet ─────────────────────────────────────────────
    summary_ws.merge_cells("A1:C1")
    title = summary_ws["A1"]
    title.value     = "✨ Misbah's Magic — Reconciliation Report"
    title.font      = Font(bold=True, size=14, color=_C["header_bg"], name="Calibri")
    title.alignment = Alignment(horizontal="center")
    summary_ws.row_dimensions[1].height = 32

    s = result.get("summary", {})
    rows = [
        ("Total Bank Transactions",        s.get("total_bank", 0),             ""),
        ("Total QuickBooks Transactions",  s.get("total_qb", 0),               ""),
        ("",                               "",                                  ""),
        ("✅  Matched",                    s.get("matched_count", 0),          _C["matched"]),
        ("❌  Missing in QuickBooks",      s.get("missing_in_qb_count", 0),    _C["missing"]),
        ("❌  Missing in Bank",            s.get("missing_in_bank_count", 0),  _C["missing"]),
        ("⚠️   Duplicates",               s.get("duplicates_count", 0),       _C["duplicate"]),
        ("🔄  Discrepancies",             s.get("discrepancies_count", 0),    _C["discrepancy"]),
    ]
    for r, (label, val, color) in enumerate(rows, 3):
        c1 = summary_ws.cell(row=r, column=1, value=label)
        c2 = summary_ws.cell(row=r, column=2, value=val)
        c1.font = _BOLD
        if color:
            c1.fill = _fill(color)
            c2.fill = _fill(color)
    _autofit(summary_ws)

    # ── Matched ───────────────────────────────────────────────────
    _add_sheet(
        wb, "✅ Matched",
        ["Bank Description", "QB Description", "Bank Date", "QB Date", "Bank Amount", "QB Amount", "Similarity %"],
        result.get("matched", []),
        lambda r: [r["bank_description"], r["qb_description"], r["bank_date"], r["qb_date"],
                   r["bank_amount"], r["qb_amount"], f'{r["similarity"]}%'],
        _C["matched"],
    )

    # ── Missing in QB ─────────────────────────────────────────────
    _add_sheet(
        wb, "❌ Missing in QB",
        ["Description", "Date", "Amount"],
        result.get("missing_in_qb", []),
        lambda r: [r["description"], r["date"], r["amount"]],
        _C["missing"],
    )

    # ── Missing in Bank ───────────────────────────────────────────
    _add_sheet(
        wb, "❌ Missing in Bank",
        ["Description", "Date", "Amount"],
        result.get("missing_in_bank", []),
        lambda r: [r["description"], r["date"], r["amount"]],
        _C["missing"],
    )

    # ── Discrepancies ─────────────────────────────────────────────
    _add_sheet(
        wb, "🔄 Discrepancies",
        ["Bank Description", "QB Description", "Bank Date", "QB Date", "Bank Amount", "QB Amount", "Difference"],
        result.get("discrepancies", []),
        lambda r: [r["bank_description"], r["qb_description"], r["bank_date"], r["qb_date"],
                   r["bank_amount"], r["qb_amount"], r["difference"]],
        _C["discrepancy"],
    )

    # ── Duplicates ────────────────────────────────────────────────
    _add_sheet(
        wb, "⚠️ Duplicates",
        ["Description", "Date", "Amount"],
        result.get("duplicates", []),
        lambda r: [r["description"], r["date"], r["amount"]],
        _C["duplicate"],
    )

    filepath = os.path.join(output_dir, f"report_{uuid.uuid4().hex[:8]}.xlsx")
    wb.save(filepath)
    return filepath
