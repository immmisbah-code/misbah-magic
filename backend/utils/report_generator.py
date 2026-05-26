"""
Report Generator
Generates a color-coded Excel reconciliation report using openpyxl.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import uuid


# Color palette
GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")  # Matched
RED_FILL    = PatternFill("solid", fgColor="FFC7CE")  # Missing
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")  # Discrepancies
ORANGE_FILL = PatternFill("solid", fgColor="FFDAAA")  # Duplicates
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")  # Header bg
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD_FONT   = Font(bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)


def _write_header(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 20


def _write_row(ws, data, row, fill=None):
    for col, val in enumerate(data, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.alignment = Alignment(vertical='center')
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)


def generate_report(result: dict, output_dir: str = "/tmp") -> str:
    """
    Generate a full Excel reconciliation report.
    Returns the file path of the saved report.
    """
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"

    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = "Misbah's Magic — Reconciliation Report"
    title_cell.font = Font(bold=True, size=14, color="1F4E79")
    title_cell.alignment = Alignment(horizontal='center')
    ws_summary.row_dimensions[1].height = 30

    summary = result.get("summary", {})
    summary_data = [
        ("Total Bank Transactions",       summary.get("total_bank", 0)),
        ("Total QuickBooks Transactions",  summary.get("total_qb", 0)),
        ("✅ Matched",                     summary.get("matched_count", 0)),
        ("❌ Missing in QuickBooks",       summary.get("missing_in_qb_count", 0)),
        ("❌ Missing in Bank",             summary.get("missing_in_bank_count", 0)),
        ("⚠️  Duplicates",                 summary.get("duplicates_count", 0)),
        ("🔄 Discrepancies",              summary.get("discrepancies_count", 0)),
    ]

    for r, (label, val) in enumerate(summary_data, 3):
        ws_summary.cell(row=r, column=1, value=label).font = BOLD_FONT
        ws_summary.cell(row=r, column=2, value=val)
    _autofit_columns(ws_summary)

    # ── Sheet 2: Matched ─────────────────────────────────────────
    ws_matched = wb.create_sheet("✅ Matched")
    headers = ["Bank Description", "QB Description", "Bank Date", "QB Date", "Bank Amount", "QB Amount", "Similarity %"]
    _write_header(ws_matched, headers)
    for r, item in enumerate(result.get("matched", []), 2):
        _write_row(ws_matched, [
            item["bank_description"], item["qb_description"],
            item["bank_date"], item["qb_date"],
            item["bank_amount"], item["qb_amount"],
            f'{item["similarity"]}%'
        ], r, fill=GREEN_FILL)
    _autofit_columns(ws_matched)

    # ── Sheet 3: Missing in QB ────────────────────────────────────
    ws_missing_qb = wb.create_sheet("❌ Missing in QB")
    _write_header(ws_missing_qb, ["Description", "Date", "Amount"])
    for r, item in enumerate(result.get("missing_in_qb", []), 2):
        _write_row(ws_missing_qb, [item["description"], item["date"], item["amount"]], r, fill=RED_FILL)
    _autofit_columns(ws_missing_qb)

    # ── Sheet 4: Missing in Bank ──────────────────────────────────
    ws_missing_bank = wb.create_sheet("❌ Missing in Bank")
    _write_header(ws_missing_bank, ["Description", "Date", "Amount"])
    for r, item in enumerate(result.get("missing_in_bank", []), 2):
        _write_row(ws_missing_bank, [item["description"], item["date"], item["amount"]], r, fill=RED_FILL)
    _autofit_columns(ws_missing_bank)

    # ── Sheet 5: Discrepancies ────────────────────────────────────
    ws_disc = wb.create_sheet("🔄 Discrepancies")
    _write_header(ws_disc, ["Bank Description", "QB Description", "Bank Date", "QB Date", "Bank Amount", "QB Amount", "Difference"])
    for r, item in enumerate(result.get("discrepancies", []), 2):
        _write_row(ws_disc, [
            item["bank_description"], item["qb_description"],
            item["bank_date"], item["qb_date"],
            item["bank_amount"], item["qb_amount"],
            item["difference"]
        ], r, fill=YELLOW_FILL)
    _autofit_columns(ws_disc)

    # ── Sheet 6: Duplicates ───────────────────────────────────────
    ws_dup = wb.create_sheet("⚠️ Duplicates")
    _write_header(ws_dup, ["Description", "Date", "Amount"])
    for r, item in enumerate(result.get("duplicates", []), 2):
        _write_row(ws_dup, [item["description"], item["date"], item["amount"]], r, fill=ORANGE_FILL)
    _autofit_columns(ws_dup)

    # Save
    filename = f"reconciliation_report_{uuid.uuid4().hex[:8]}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    return filepath
